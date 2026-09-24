#!/usr/bin/env python3
"""Interactive Nutanix cluster utilization report (CPU, memory, storage) via Prism Central v4.0 APIs.

Made by Guransh Deol.
"""

import getpass
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
import urllib3
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CLUSTERS_URL = "/api/clustermgmt/v4.0/config/clusters"
STATS_URL = "/api/clustermgmt/v4.0/stats/clusters/{ext_id}"
STAT_FIELDS = [
    "hypervisorCpuUsagePpm",
    "aggregateHypervisorMemoryUsagePpm",
    "storageUsageBytes",
    "storageCapacityBytes",
]
PAGE_LIMIT = 100
GIB = 1024 ** 3

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill("solid", fgColor="EEF3FA")
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def prompt(text, default=None):
    suffix = f" [{default}]" if default is not None else ""
    value = input(f"{text}{suffix}: ").strip()
    return value or (str(default) if default is not None else "")


def prompt_yes_no(text, default=False):
    hint = "Y/n" if default else "y/N"
    value = input(f"{text} ({hint}): ").strip().lower()
    return default if not value else value in ("y", "yes")


class PrismCentral:
    def __init__(self, host, port, username, password, verify_ssl):
        self.base = f"https://{host}:{port}"
        self.session = requests.Session()
        self.session.auth = (username, password)
        self.session.verify = verify_ssl
        self.session.headers.update({"Accept": "application/json"})

    def get(self, path, params=None):
        resp = self.session.get(self.base + path, params=params, timeout=60)
        if resp.status_code == 401:
            raise PermissionError("Authentication failed (401). Check username/password.")
        resp.raise_for_status()
        return resp.json()

    def list_clusters(self):
        clusters, page = [], 0
        while True:
            body = self.get(CLUSTERS_URL, {"$page": page, "$limit": PAGE_LIMIT})
            data = body.get("data") or []
            clusters.extend(data)
            total = (body.get("metadata") or {}).get("totalAvailableResults", len(clusters))
            if not data or len(clusters) >= total:
                break
            page += 1
        # The Prism Central instance itself is listed as a cluster; it has no workload stats.
        return [
            c for c in clusters
            if "PRISM_CENTRAL" not in ((c.get("config") or {}).get("clusterFunction") or [])
        ]

    def cluster_stats(self, ext_id, start, end, stat_type):
        params = {
            "$startTime": start,
            "$endTime": end,
            "$samplingInterval": 86400,
            "$statType": stat_type,
            "$select": ",".join(STAT_FIELDS),
        }
        return self.get(STATS_URL.format(ext_id=ext_id), params).get("data") or {}


def series_by_day(stats, field):
    out = {}
    for point in stats.get(field) or []:
        if point.get("value") is None or not point.get("timestamp"):
            continue
        day = datetime.fromisoformat(point["timestamp"].replace("Z", "+00:00")).date()
        out[day] = point["value"]
    return out


def collect(pc, cluster, start, end):
    avg = pc.cluster_stats(cluster["extId"], start, end, "AVG")
    peak = pc.cluster_stats(cluster["extId"], start, end, "MAX")

    cpu_avg = series_by_day(avg, "hypervisorCpuUsagePpm")
    cpu_max = series_by_day(peak, "hypervisorCpuUsagePpm")
    mem_avg = series_by_day(avg, "aggregateHypervisorMemoryUsagePpm")
    mem_max = series_by_day(peak, "aggregateHypervisorMemoryUsagePpm")
    used = series_by_day(avg, "storageUsageBytes")
    used_max = series_by_day(peak, "storageUsageBytes")
    cap = series_by_day(avg, "storageCapacityBytes")

    days = sorted(set(cpu_avg) | set(mem_avg) | set(used))
    rows = []
    for d in days:
        capacity = cap.get(d)
        storage_used = used.get(d)
        rows.append({
            "date": d,
            "cpu_avg": ppm_to_pct(cpu_avg.get(d)),
            "cpu_max": ppm_to_pct(cpu_max.get(d)),
            "mem_avg": ppm_to_pct(mem_avg.get(d)),
            "mem_max": ppm_to_pct(mem_max.get(d)),
            "storage_used_gib": to_gib(storage_used),
            "storage_max_gib": to_gib(used_max.get(d)),
            "storage_cap_gib": to_gib(capacity),
            "storage_pct": round(storage_used / capacity * 100, 2) if storage_used and capacity else None,
        })
    return rows


def ppm_to_pct(value):
    return round(value / 10_000, 2) if value is not None else None


def to_gib(value):
    return round(value / GIB, 2) if value is not None else None


def stats_of(rows, key, fn):
    values = [r[key] for r in rows if r[key] is not None]
    return round(fn(values), 2) if values else None


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font, cell.alignment, cell.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
    ws.row_dimensions[row].height = 32


def style_body(ws, first_row, last_row, ncols):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            if (r - first_row) % 2:
                cell.fill = ZEBRA_FILL


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_heatmap(ws, cell_range):
    ws.conditional_formatting.add(cell_range, ColorScaleRule(
        start_type="num", start_value=0, start_color="63BE7B",
        mid_type="num", mid_value=60, mid_color="FFEB84",
        end_type="num", end_value=100, end_color="F8696B",
    ))


def write_summary(wb, results, pc_host, start, end):
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Nutanix Cluster Utilization Report - Last 30 Days"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Prism Central: {pc_host}    Period (UTC): {start[:10]} to {end[:10]}    Generated: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(italic=True, color="595959")

    headers = [
        "Cluster", "Avg CPU %", "Peak CPU %", "Avg Memory %", "Peak Memory %",
        "Storage Used (GiB)", "Storage Capacity (GiB)", "Storage Used %", "Storage Growth (GiB)", "Days of Data",
    ]
    hr = 4
    for i, h in enumerate(headers, 1):
        ws.cell(row=hr, column=i, value=h)
    style_header(ws, hr, len(headers))

    r = hr
    for name, rows in results.items():
        r += 1
        latest = next((x for x in reversed(rows) if x["storage_used_gib"] is not None), None)
        first = next((x for x in rows if x["storage_used_gib"] is not None), None)
        ws.append([
            name,
            stats_of(rows, "cpu_avg", lambda v: sum(v) / len(v)),
            stats_of(rows, "cpu_max", max),
            stats_of(rows, "mem_avg", lambda v: sum(v) / len(v)),
            stats_of(rows, "mem_max", max),
            latest["storage_used_gib"] if latest else None,
            latest["storage_cap_gib"] if latest else None,
            latest["storage_pct"] if latest else None,
            round(latest["storage_used_gib"] - first["storage_used_gib"], 2) if latest and first else None,
            len(rows),
        ])
    if r > hr:
        style_body(ws, hr + 1, r, len(headers))
        for col in "BCDEH":
            add_heatmap(ws, f"{col}{hr + 1}:{col}{r}")
        for row in ws.iter_rows(min_row=hr + 1, max_row=r, min_col=2, max_col=9):
            for cell in row:
                cell.number_format = "#,##0.00"
        ws.cell(row=hr + 1, column=1).alignment = Alignment(horizontal="left")
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    set_widths(ws, [30, 12, 12, 14, 14, 18, 20, 15, 20, 13])


def write_cluster_sheet(wb, name, rows):
    title = "".join(ch for ch in name if ch not in '[]:*?/\\')[:31] or "Cluster"
    while title in wb.sheetnames:
        title = title[:28] + f"_{len(wb.sheetnames)}"
    ws = wb.create_sheet(title)
    ws["A1"] = f"Daily Utilization - {name}"
    ws["A1"].font = TITLE_FONT

    headers = [
        "Date", "Avg CPU %", "Peak CPU %", "Avg Memory %", "Peak Memory %",
        "Storage Used (GiB)", "Peak Storage Used (GiB)", "Storage Capacity (GiB)", "Storage Used %",
    ]
    hr = 3
    for i, h in enumerate(headers, 1):
        ws.cell(row=hr, column=i, value=h)
    style_header(ws, hr, len(headers))

    for row in rows:
        ws.append([row["date"], row["cpu_avg"], row["cpu_max"], row["mem_avg"], row["mem_max"],
                   row["storage_used_gib"], row["storage_max_gib"], row["storage_cap_gib"], row["storage_pct"]])
    last = hr + len(rows)
    set_widths(ws, [13, 12, 12, 14, 14, 18, 22, 22, 15])
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    if not rows:
        ws.cell(row=hr + 1, column=1, value="No stats returned for this period.")
        return

    style_body(ws, hr + 1, last, len(headers))
    for r in range(hr + 1, last + 1):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        for c in range(2, 10):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    for col in "BCDEI":
        add_heatmap(ws, f"{col}{hr + 1}:{col}{last}")

    dates = Reference(ws, min_col=1, min_row=hr + 1, max_row=last)
    anchor_col = get_column_letter(len(headers) + 2)
    charts = [
        ("CPU Usage %", "Percent", [2, 3], 0, 100),
        ("Memory Usage %", "Percent", [4, 5], 0, 100),
        ("Storage Used vs Capacity (GiB)", "GiB", [6, 8], None, None),
    ]
    for idx, (chart_title, y_title, cols, y_min, y_max) in enumerate(charts):
        chart = LineChart()
        chart.title, chart.y_axis.title, chart.x_axis.title = chart_title, y_title, "Date"
        chart.height, chart.width = 7.5, 18
        chart.x_axis.number_format = "mm-dd"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        if y_min is not None:
            chart.y_axis.scaling.min, chart.y_axis.scaling.max = y_min, y_max
        for c in cols:
            chart.add_data(Reference(ws, min_col=c, min_row=hr, max_row=last), titles_from_data=True)
        chart.set_categories(dates)
        ws.add_chart(chart, f"{anchor_col}{hr + idx * 16}")


def main():
    print("=" * 60)
    print(" Nutanix Cluster Utilization Report (last 30 days)")
    print("=" * 60)
    host = prompt("Prism Central IP/FQDN")
    if not host:
        sys.exit("Prism Central address is required.")
    port = prompt("Port", 9440)
    username = prompt("Local admin username", "admin")
    password = getpass.getpass("Password: ")
    verify_ssl = prompt_yes_no("Verify SSL certificate?", default=False)
    if not verify_ssl:
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    pc = PrismCentral(host, port, username, password, verify_ssl)
    try:
        clusters = pc.list_clusters()
    except PermissionError as e:
        sys.exit(str(e))
    except requests.RequestException as e:
        sys.exit(f"Could not reach Prism Central: {e}")
    if not clusters:
        sys.exit("No clusters registered with this Prism Central.")

    print(f"\nFound {len(clusters)} cluster(s):")
    for i, c in enumerate(clusters, 1):
        print(f"  {i}. {c.get('name')}")
    choice = prompt("\nClusters to include (e.g. 1,3) or 'all'", "all")
    if choice.lower() != "all":
        try:
            picks = {int(x) for x in choice.replace(" ", "").split(",") if x}
            clusters = [c for i, c in enumerate(clusters, 1) if i in picks]
        except ValueError:
            sys.exit("Invalid selection.")
        if not clusters:
            sys.exit("No valid clusters selected.")

    end_dt = datetime.now(timezone.utc).replace(microsecond=0)
    start_dt = end_dt - timedelta(days=30)
    start, end = start_dt.isoformat(), end_dt.isoformat()

    results = {}
    for c in clusters:
        name = c.get("name") or c["extId"]
        print(f"Fetching stats for {name} ...", end=" ", flush=True)
        try:
            results[name] = collect(pc, c, start, end)
            print(f"{len(results[name])} day(s)")
        except requests.RequestException as e:
            print(f"FAILED ({e})")
            results[name] = []

    default_name = f"cluster_usage_{end_dt:%Y%m%d_%H%M}.xlsx"
    out = Path(prompt("\nOutput file", default_name)).expanduser()
    if out.suffix.lower() != ".xlsx":
        out = out.with_suffix(".xlsx")

    wb = Workbook()
    write_summary(wb, results, host, start, end)
    for name, rows in results.items():
        write_cluster_sheet(wb, name, rows)
    wb.save(out)
    print(f"\nReport saved to {out.resolve()}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("\nCancelled.")
